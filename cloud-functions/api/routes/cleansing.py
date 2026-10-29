"""原生清洗路由(方案 B): 文件识别/预览/执行/模板(6 类目标)

契约要点(CleansingPage 直连, 平铺响应——axios 拦截器对无 data 字段的 {ok,...} 保留原样):
- detect → {ok, columns:[{name}], total}(平铺, 不能 {ok,data} 包装否则前端 d.ok undefined)
- preview → {ok, preview:[{target字段}], total}(preview 为清洗后前 50 行)
- execute-async → {ok, task_id, success, failed, error?, message?}(同步执行完返回, 前端优先生成 success 直显)
- templates: GET 标准 {ok,data:[...]}; POST {name,doc_type,mapping} → {ok,data:{message}}
"""
import csv
import io
import json
import re
import time
from datetime import datetime, timedelta, timezone  # datetime 用于 status 超时标记(缺 datetime 曾 NameError——同步分支/轮询全崩)

from fastapi import APIRouter, Query
from fastapi import Request
from fastapi import UploadFile
from fastapi import File
from fastapi import Form

from db import query, one, execute, executemany
from routes.cleanse_write import _finalize_import, _adjust_inventory, _write_rows, _write_batch, _sync_batches
from routes.common import ok, fail, traced, try_err

router = APIRouter(tags=["cleansing"])

# 目标类型 → 表/仓库类型映射
_TARGET_WH = {"inventory": "own", "platform_inv": "platform", "inventory_b": "platform_b"}
_NUMS = {"number", "price", "quantity", "score", "safety_qty", "available_qty", "in_transit_qty",
         "month_inbound", "month_outbound", "beginning_stock", "turnover_days", "c_transit",
         "locked_qty", "weight", "volume", "box_qty", "total_amount", "unit_price",
         "freight_amount", "subsidy_amount", "tax_amount", "discount_amount", "actual_amount"}


def _mk_upload(data: bytes, filename: str = "file.csv"):
    """构造 UploadFile(raw body/ base64 路径——绕 iOS multipart body 空)"""
    import io
    from starlette.datastructures import UploadFile as _UF, Headers as _H
    # 魔数嗅探: xlsx(zip) PK 头 → filename 置 xlsx——否则退化 csv 当 CSV 解析乱码单列
    if data[:4] == b"PK\x03\x04":
        filename = "file.xlsx"
    return _UF(io.BytesIO(data), filename=filename,
               headers=_H({"content-type": "application/octet-stream"}), size=len(data))


def _read_file_bytes(upload: UploadFile):
    data = upload.file.read()
    name = (upload.filename or "upload.csv").lower()
    return data, name


def _parse_table(data: bytes, fname: str):
    """解析 CSV/XLSX → (headers, rows); rows 为 list[dict]"""
    if fname.endswith(".xlsx"):
        import openpyxl
        # 小文件普通模式(兼容非标准/无维度信息文件): read_only 下 max_row/max_column 可能误判 1×1
        # (实测: 采购订单明细导出(非图书).xlsx read_only 读到 1列0行, 普通模式 48列9行)
        # 大文件(>=10MB)仍用 read_only(标准文件可靠, 省内存)
        _rb = len(data) >= 10 * 1024 * 1024
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=_rb, data_only=True)
        except Exception:
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        # 取第一个有数据的 sheet(active 可能是空壳/标题页)
        _ws = None
        for _s in wb.worksheets:
            if _s.max_row and _s.max_column:
                _ws = _s
                break
        if _ws is None:
            _ws = wb.active
        rows_iter = _ws.iter_rows(values_only=True)
        headers = None
        rows = []
        for r in rows_iter:
            if headers is None:
                headers = [str(c).strip() if c is not None else "" for c in r]
                continue
            rows.append(dict(zip(headers, ["" if c is None else c for c in r])))
        headers = [h for h in headers if h]
        return headers, rows
    # CSV: 自动探测编码/BOM/分隔符
    raw = data
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except Exception:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    text = text.replace("\x00", "")  # NUL 清洗
    text = text.replace("\r\n", "\n").replace("\r", "\n")  # 2026-10-14 换行归一(\r 老Mac——csv universal newline 防 'new-line character seen')
    try:
        # 分隔符探测扩展到多平台导出常见分隔: 逗号/制表/分号/竖线/中文逗号顿号
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|，、")
        sep = dialect.delimiter
    except Exception:
        sep = ","
    reader = csv.reader(io.StringIO(text), delimiter=sep)
    all_rows = list(reader)
    # 单列/误判兜底: Sniffer 对竖线/中文逗号等常失效(实测), 手动探测候选分隔符
    # 判定: 候选分隔后表头行分裂 >1 列且前 5 行列数一致 → 采用
    if len(all_rows) <= 1 or not (all_rows and max((len(r) for r in all_rows[:10]), default=1) > 1):
        for _s in ("\t", ";", "|", "，", "、", ","):
            _rr = list(csv.reader(io.StringIO(text), delimiter=_s))
            if _rr and len(_rr[0]) > 1 and \
               all(len(r) == len(_rr[0]) for r in _rr[:5] if any(str(x).strip() for x in r)):
                sep, all_rows = _s, _rr
                break
    if not all_rows:
        return [], []
    headers = [h.strip() for h in all_rows[0] if h.strip()]
    rows = []
    for r in all_rows[1:]:
        if not any(str(c).strip() for c in r):
            continue
        rows.append(dict(zip(headers, r[:len(headers)])))
    return headers, rows


def _v(value, vtype):
    """按目标类型清洗单值"""
    if value is None:
        return ""
    s = str(value).strip()
    if vtype == "number":
        s = re.sub(r"[,\s元¥￥]", "", s)
        try:
            return float(s)
        except Exception:
            return 0.0
    if vtype == "date":
        m = re.match(r"(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})", s)
        if m:
            return "%s-%02d-%02d" % (m.group(1), int(m.group(2)), int(m.group(3)))
        return s[:10]
    return s


def _clean_rows(rows, mapping, custom_fields=None):
    """rows → 按 mapping {src:{target,type}} 清洗; 忽略无映射列; 空 target 跳过"""
    out = []
    for r in rows:
        item = {}
        for src, cfg in (mapping or {}).items():
            if src == "_meta" or not cfg or not cfg.get("target"):
                continue
            v = r.get(src)
            vtype = cfg.get("type") or "string"
            item[cfg["target"]] = _v(v, vtype)
        out.append(item)
    return out


def _after(now, days_ago):
    return (now - timedelta(days=days_ago)).strftime("%Y-%m-%d")


def _detect_columns(fname, data):
    headers, _ = _parse_table(data, fname)
    return headers


# ── detect: 识别列名 ────────────────────────────────────────────────────
@router.post("/cleansing/dup-check")
@traced
async def cleansing_dup_check(request: Request, file: UploadFile = File(None),
                              q_target: str = Query(""), q_channel: str = Query(""),
                              q_mp: str = Query("")):
    """2026-10-17 前端提前检测专用(职责: 前端拦截/后端写入): 三键优先(order_no|sku|barcode,
    空 barcode 落二键 order_no|sku)查库已存在——纯只读 SELECT 秒级, 不写入不建任务不拖导入性能
    返回 {duplicate:n(全重), overlap:n(部分), new:n}; 前端据此拦截/弹窗 force_new"""
    try:
        _t = q_target or "order"
        if _t not in ("order", "gmv_flow"):
            return {"ok": True, "duplicate": 0, "overlap": 0, "new": 0}
        if not file:
            _raw = await request.body()
            if _raw:
                file = _mk_upload(_raw, "file.csv")
        data, fname = _read_file_bytes(file)
        _ch = q_channel or "jd"
        _, _rows = _parse_table(data, fname)
        if not _rows:
            return {"ok": True, "duplicate": 0, "overlap": 0, "new": 0}
        _mp = json.loads(q_mp or "{}")
        _cb = _clean_rows(_rows, _mp)
        if _t == "order":
            _keys = [
                (str(r.get("order_no") or "") + "|" + str(r.get("sku") or "")
                 + (("|" + str(r.get("barcode") or "").strip()) if str(r.get("barcode") or "").strip() else ""))
                for r in _cb if r.get("order_no") and r.get("sku")]
            _expr = "CONCAT_WS('|', `order_no`, `sku`, NULLIF(`barcode`,''))"
            _tbl, _soft = "orders", " AND (deleted_at IS NULL OR deleted_at='')"
        else:
            _keys = [str(r.get("platform") or "") + "|" + str(r.get("flow_no") or "") for r in _cb
                     if r.get("platform") and r.get("flow_no")]
            _expr = "CONCAT(`platform`,'|',`flow_no`)"
            _tbl, _soft = "gmv_flows", ""
        _uk = list(dict.fromkeys(_keys))
        if not _uk:
            return {"ok": True, "duplicate": 0, "overlap": 0, "new": 0}
        _exists = 0
        for _i in range(0, len(_uk), 200):
            _batch = _uk[_i:_i + 200]
            _ph = ",".join(["%s"] * len(_batch))
            _cnt = one("SELECT COUNT(*) c FROM `%s` WHERE channel=%%s%s "
                       "AND %s IN (%s)" % (_tbl, _soft, _expr, _ph),
                       [_ch] + _batch)
            _exists += int((_cnt or {}).get("c") or 0)
        _n = len(set(_keys))
        return {"ok": True, "duplicate": 1 if _exists >= _n else 0,
                "overlap": _exists if 0 < _exists < _n else 0,
                "new": _n - _exists}
    except Exception as _e:
        from routes.common import try_err as _terr
        _terr('cleansing', 'dup-check 降级(放行由后端二键兜底)', _e)
        return {"ok": True, "duplicate": 0, "overlap": 0, "new": 0}

@router.post("/cleansing/driver")
@traced
def cleansing_driver(task_id: str = "", driver: str = "", channel: str = ""):
    """前端轮询驱动标记(2026-10-15): 前端轮询时写 params.driver=frontend——TaskCard 显示推进方式(前端轮询/后端推进)"""
    if task_id and driver in ("frontend", "backend"):
        try:
            execute("UPDATE sync_tasks SET params=JSON_SET(params, '$.driver', %s) WHERE task_id=%s", [driver, task_id])
        except Exception:
            try:
                _row = one("SELECT params FROM sync_tasks WHERE task_id=%s", [task_id])
                _pj = json.loads((_row or {}).get("params") or "{}")
                _pj["driver"] = driver
                execute("UPDATE sync_tasks SET params=%s WHERE task_id=%s", [json.dumps(_pj, ensure_ascii=False), task_id])
            except Exception as _e:
                    try_err('cleansing', 'driver 标记降级', _e)
    return ok({})

@router.post("/cleansing/detect")
@traced
async def cleansing_detect(request: Request, file: UploadFile = File(None)):
    if not file:
        _raw = await request.body()
        if _raw:
            file = _mk_upload(_raw, "file.xlsx")
    # 临时诊断(2026-10-14 422 file missing 定位): 记录请求表单/文件到达情况
    try:
        from db import execute as _ed
        _ed("INSERT INTO quality_logs(log_type, level, message, source) VALUES('perf_debug','info',%s,'cleansing')",
            ("detect 收到 file: %s, filename: %s, size: %s" % (
                bool(file), (file.filename or '')[:40], getattr(file, 'size', None))))
    except Exception:
        pass
    data, fname = _read_file_bytes(file)
    try:
        headers, rows = _parse_table(data, fname)
        return {"ok": True, "columns": [{"name": h} for h in headers], "total": len(rows)}
    except Exception as e:
        return {"ok": False, "error": "文件解析失败: %s" % str(e)[:200]}


# ── preview: 清洗预览(前 50 行) ─────────────────────────────────────────
@router.post("/cleansing/preview")
@traced
async def cleansing_preview(request: Request, file: UploadFile = File(None),
                            q_target: str = Query(""), q_channel: str = Query(""), q_conflict: str = Query(""), q_mp: str = Query("")):
    # 新方式(2026-10-14 根治): raw body 直传(绕 iOS multipart body 空) + 参数走 query
    if not file:
        _raw = await request.body()
        if _raw:
            file = _mk_upload(_raw, "file.csv")
    target = q_target or "order"
    channel = q_channel or "jd"
    conflict_mode = q_conflict or "sum"
    mapping = q_mp or "{}"
    data, fname = _read_file_bytes(file)
    try:
        _, rows = _parse_table(data, fname)
        mp = json.loads(mapping or "{}")
        cleaned = _clean_rows(rows, mp)
        # 零映射兜底(严谨性): 无任何有效值 → 明确提示而非静默返回 0 条
        if not any(any(str(v) not in ("", "None", "none") for v in it.values()) for it in cleaned):
            return {"ok": False, "error": "未配置有效列映射，请先在字段映射界面选择目标字段后再预览"}
        return {"ok": True, "preview": cleaned[:50], "total": len(cleaned),
                "target": target, "channel": channel, "conflict_mode": conflict_mode}
    except Exception as e:
        return {"ok": False, "error": "预览失败: %s" % str(e)[:200]}


# ── execute-async: 全量清洗写入 ─────────────────────────────────────────
@router.post("/cleansing/execute-async")
@traced
async def cleansing_execute(request: Request, file: UploadFile = File(None),
                            q_target: str = Query(""), q_channel: str = Query(""), q_conflict: str = Query(""), q_mp: str = Query("")):
    if not file:
        _raw = await request.body()
        if _raw:
            file = _mk_upload(_raw, "file.csv")
    target = q_target or "order"
    channel = q_channel or "jd"
    conflict_mode = q_conflict or "sum"
    mapping = q_mp or "{}"
    import logging as _lg
    """导入(接力式: POST 存源文件+建任务 → GET /cleansing/status 续跑分窗清洗写入,
    防大批量 120s 超时——与导出/seed 同模式; Makers 实例请求后可冻结, 不用纯后台线程)
    返回 {ok, task_id}"""
    # 2026-10-14 防重: 同渠道已有进行中导入任务(误触/重复提交同批) → 拒绝
    try:
        _busy = one("SELECT COUNT(*) AS c FROM sync_tasks WHERE task_type='cleansing' "
                    "AND status='running' AND channel=%s AND updated_at > NOW() - INTERVAL 30 MINUTE", [channel]) or {}
        if int(_busy.get("c") or 0) > 0:
            return {"ok": False, "error": "已有进行中的导入任务，请等待完成后再提交（防止同批重复导入）"}
    except Exception:
        pass
    data, fname = _read_file_bytes(file)
    try_err('cleansing', 'execute target=%s fname=%s' % (target, fname))
    # 2026-10-15 重复提前检测: 同键已存在(全部重复——二次导入同表)→ 同步返回跳过计数(不建任务——秒级拦截——不等异步 4 分钟)
    # 2026-10-16 泛化: order(键=order_no|sku) + gmv_flow(键=platform|flow_no) 对称拦截
    #   —— gmv_flow 此前无提前检测(仅 DB 唯一索引兜底, 二次导入全批 1062 失败无因)——用户实测重复导入痛点
    # 2026-10-17 职责边界(用户): 前端负责检测拦截(三键优先 order_no+sku+barcode / 二键兜底),
    #   后端只做二键写入检测(保持性能稳定, 不逐行算三键)
    try:
        from db import query as _qk
        _dr = _mk_upload(data, fname)
        _bdata, _bfname = _read_file_bytes(_dr)
        try:
            _bh, _brows = _parse_table(_bdata, _bfname)
            _dup_cfg = {"order": ("orders", "order_no", "sku"),
                        "gmv_flow": ("gmv_flows", "platform", "flow_no")}
            if target in _dup_cfg and _brows:
                # 2026-10-15 修复: 原始行无目标键(中文列名)——需清洗后取键(否则 _keys 空不拦截——重复仍建任务)
                _tbl_d, _k1_d, _k2_d = _dup_cfg[target]
                _mapping = json.loads(mapping or "{}")
                _cb = _clean_rows(_brows, _mapping)
                _keys = [str(r.get(_k1_d) or "") + "|" + str(r.get(_k2_d) or "") for r in _cb
                         if r.get(_k1_d) and r.get(_k2_d)]
                _uk = list(dict.fromkeys(_keys))
                _exists = 0
                for _i in range(0, len(_uk), 200):
                    _batch = _uk[_i:_i+200]
                    _pairs = [(b.split("|")[0], b.split("|")[1]) for b in _batch]
                    # 同键 EXISTS——分片查(order 需排除软删——软删单可重导; gmv_flows 无 deleted_at 列)
                    _in_ph = ",".join(["%s"] * len(_batch))
                    _soft = (" AND (deleted_at IS NULL OR deleted_at='')" if _tbl_d == "orders" else "")
                    _sql = ("SELECT COUNT(*) c FROM `%s` WHERE channel=%%s%s "
                            "AND CONCAT(`%s`,'|',`%s`) IN (%s)" % (_tbl_d, _soft, _k1_d, _k2_d, _in_ph))
                    _cnt = one(_sql, [channel] + _batch)
                    _exists += int((_cnt or {}).get("c") or 0)
                _force_new = request.query_params.get("force_new", "")
                if _keys and not _force_new:
                    _uk_n = len(set(_keys))
                    if _exists >= _uk_n:
                        # 全部已存在——重复导入——同步返回跳过(不建任务)
                        return {"ok": True, "task_id": "", "target": target, "success": 0, "failed": len(_keys),
                                "total_rows": len(_brows), "duplicate": True,
                                "message": "文件已导入过（同键全部已存在），已跳过 %d 条" % len(_keys)}
                    if _exists > 0:
                        # 部分重叠(2026-10-16): 返回 overlap——前端弹窗"X 条已存在(新增 Y 条)——确认只写新增?"
                        return {"ok": True, "task_id": "", "success": 0, "failed": _exists,
                                "total_rows": len(_brows), "overlap_partial": True,
                                "overlap": _exists, "new": _uk_n - _exists,
                                "message": "%d 条已存在（新增 %d 条）——确认导入新增？" % (_exists, _uk_n - _exists)}
        except Exception:
            pass
    except Exception:
        pass
    task_id = "clean_%d_%04d" % (int(time.time()), int(time.time() * 1000) % 10000)
    try:
        # source 文件存 export_files(clean_ 前缀, 7 天清理覆盖)
        import base64 as _b64
        try:
            execute("INSERT INTO export_files(filename, content, channel) VALUES(%s,%s,%s) "
                    "ON DUPLICATE KEY UPDATE content=VALUES(content)",
                    ("clean_" + task_id, data, channel))
        except Exception:
            execute("INSERT INTO export_files(filename, content, channel) VALUES(%s,%s,%s) "
                    "ON DUPLICATE KEY UPDATE content=VALUES(content)",
                    ("clean_" + task_id, "base64:" + _b64.b64encode(data).decode("ascii"), channel))
        params = {"fname": fname, "mp": mapping, "target": target, "channel": channel,
                  "conflict_mode": conflict_mode, "offset": 0, "success": 0, "failed": 0,
                  "total": 0}
        execute("INSERT INTO sync_tasks(task_id, task_type, status, params, result, channel) "
                "VALUES(%s,'cleansing','running',%s,'{}',%s)",
                (task_id, json.dumps(params, ensure_ascii=False), channel))
        # 同步分支(2026-10-14): 小行数(<1000)请求内完成秒级返回——不依赖轮询续跑; >=1000 走接力
        try:
            _, rows0 = _parse_table(data, fname)
            if len(rows0) < 1000:
                # 同步分支(2026-10-14 修复): 任务刚建锁未满足(updated_at<3s 返回 running)——sleep 3 等锁; done 才 success, running 返 task_id 前端轮询
                import time as _ts
                _ts.sleep(4)  # 锁需 updated_at < NOW-3s(建任务 3s 后仍边界 false——4s 稳)
                _sr = cleansing_status(task_id, page_size=5000)
                _sd = _sr.get("data") or _sr
                if _sr.get("ok") is not False and _sd.get("status") == "done":
                    return {"ok": True, "task_id": task_id, "target": target,
                            "success": _sd.get("success", 0), "failed": _sd.get("failed", 0),
                            "total_rows": len(rows0),
                            "failed_details": _sd.get("failed_details") or []}
        except Exception:
            pass  # 解析/同步异常 → 走接力(由 status 续跑兜底)
        # 2026-10-15 异步返回 total_rows(前端显示已提交 N 条——非问号): len(rows0) 已解析可用
        _tr = 0
        try:
            _, _r0 = _parse_table(data, fname)
            _tr = len(_r0)
        except Exception:
            pass
        return {"ok": True, "task_id": task_id, "total_rows": _tr}
    except Exception as e:
        import traceback as _tb
        try:
            execute("INSERT INTO sync_tasks(task_id, task_type, status, params, result, channel) "
                    "VALUES(%s,'cleansing','error',%s,%s,%s)",
                    (task_id, '{}',
                     json.dumps({"error": str(e)[:400], "tb": _tb.format_exc()[-1200:]},
                                ensure_ascii=False), channel))
        except Exception:
            pass
        return fail("导入启动失败: %s" % str(e)[:200])




@router.get("/cleansing/status")
@traced
def cleansing_status(task_id: str = "", page_size: int = 5000):
    """导入续跑(前端轮询触发): 每步读源文件 → 解析 → 窗口清洗写入(5000 行/步, ≤120s)
    全部完成 → 最后步全量重清洗做库存联动+规则评估+汇总 → done"""
    row = one("SELECT status, params, result, updated_at FROM sync_tasks WHERE task_id=%s", [task_id])
    if not row:
        return {"ok": False, "error": "任务不存在"}
    status = row.get("status") or ""
    if status == "running":
        # 2026-10-14 超时自愈: running 超过 30 分钟无更新(前端轮询停/实例冻结)——标记 error 防永久卡(释放防重拦截)
        ua = str(row.get("updated_at") or "")[:19]
        if ua and ua < (datetime.now(timezone.utc) - timedelta(minutes=30)).strftime("%Y-%m-%d %H:%M:%S"):
            execute("UPDATE sync_tasks SET status='error', result=%s WHERE task_id=%s",
                    [json.dumps({"result": {"error": "任务超时(30 分钟无更新)——已自动标记结束，请重新导入"}}), task_id])
            row["status"] = "error"
    if status in ("done", "error"):
        _ri = (json.loads(row.get("result") or "{}") or {}).get("result") or {}
        return {"ok": True, "task_id": task_id, "status": status,
                "success": _ri.get("success"), "failed": _ri.get("failed"),
                "message": _ri.get("message"), "error": _ri.get("error"),
                "failed_details": _ri.get("failed_details") or []}
    # 抢占锁(防并发续跑): 锁 3s(窗间快——done 步防重跑用 finalizing 标记——见 done 分支)
    try:
        _locked = execute("UPDATE sync_tasks SET updated_at=NOW() WHERE task_id=%s "
                          "AND (updated_at IS NULL OR updated_at < DATE_SUB(NOW(), INTERVAL 3 SECOND))",
                          [task_id])
        if not _locked:
            return {"ok": True, "task_id": task_id, "status": "running"}
    except Exception as _e:
            try_err('cleansing', '静默降级', _e)
    import time as _t
    started = _t.time()
    try:
        params = json.loads(row.get("params") or "{}")
        fname = params.get("fname") or "import.csv"
        # 2026-10-16 done 步中(finalizing=1)→ 返回 running 不重跑(锁已释放但 done 步进行中——防重跑覆盖计数)
        # 超时恢复: finalizing>5分钟且无更新(updated_at 旧——done 步请求被 Makers 冻结/进程死——连接死 TiDB 自动回滚未落库)
        # → 清标记允许重跑 done 步(事务幂等——重跑干净——根治进程死永久卡)
        try:
            if params.get("finalizing"):
                _fs = one("SELECT updated_at FROM sync_tasks WHERE task_id=%s", [task_id])
                _fu = str((_fs or {}).get("updated_at") or "")[:19].replace("T", " ")
                if _fu and _fu < (datetime.now(timezone.utc) - timedelta(minutes=5)).strftime("%Y-%m-%d %H:%M:%S"):
                    execute("UPDATE sync_tasks SET params=%s, updated_at=NOW() WHERE task_id=%s",
                            [json.dumps({**params, "finalizing": 0}, ensure_ascii=False), task_id])
                    try_err('cleansing', 'finalizing 超时恢复(>5min——done 步死连接已回滚——重跑)', '')
                else:
                    return {"ok": True, "task_id": task_id, "status": "running"}
        except Exception:
            pass
        mp = json.loads(params.get("mp") or "{}")
        target = params.get("target") or "order"
        channel = params.get("channel") or "jd"
        conflict_mode = params.get("conflict_mode") or "sum"
        offset = int(params.get("offset") or 0)
        success = int(params.get("success") or 0)
        failed = int(params.get("failed") or 0)
        # 读源文件(局部 import base64——覆盖模块级 def _b64(line 399)冲突——曾 _b64.b64decode attr error 任务 error)
        import base64 as _b64x
        fr = one("SELECT content FROM export_files WHERE filename=%s", ["clean_" + task_id])
        c = fr.get("content") if fr else None
        if isinstance(c, str):
            data = _b64x.b64decode(c[7:]) if c.startswith("base64:") else c.encode("utf-8")
        elif c is not None:
            data = bytes(c)
        else:
            raise Exception("源文件不存在")
        headers, rows = _parse_table(data, fname)
        total = len(rows)
        # 分窗清洗+写入(2026-10-15 进度中: 窗口开始即设 cur_step——前端 TaskCard 实时显示进行中)
        _win_no = offset // page_size + 1
        execute("UPDATE sync_tasks SET params=%s, updated_at=NOW() WHERE task_id=%s",
                (json.dumps({**params, "cur_step": {"name": "清洗写入窗口 %d/%d" % (_win_no, max((total // page_size) + 1, 1)), "started": int(_t.time())}}, ensure_ascii=False), task_id))
        window = rows[offset:offset + page_size]
        cleaned = _clean_rows(window, mp)
        if target == "order":
            try:
                from routes.suppliers import _norm_column_value as _ncol
                for x in cleaned:
                    if x.get("order_status"):
                        x["order_status"] = _ncol(channel, "order_status", str(x["order_status"]))
            except Exception as _e:
                    try_err('cleansing', '静默降级', _e)
        # 2026-10-16 force_new 只写新增: 确认导入后过滤重叠行(unique idx 兜底防重复插入失败)
        #   —— 泛化 2026-10-16: order(键=order_no|sku) + gmv_flow(键=platform|flow_no) 对称
        # 2026-10-17 职责边界(用户): 前端检测三键优先, 后端 force_new 保持二键(写入性能稳定)
        if params.get("force_new") and target in ("order", "gmv_flow") and cleaned:
            try:
                _fk_cfg = {"order": ("orders", "order_no", "sku"),
                           "gmv_flow": ("gmv_flows", "platform", "flow_no")}
                _ftbl, _fk1, _fk2 = _fk_cfg[target]
                _fkeys = [str(r.get(_fk1) or "") + "|" + str(r.get(_fk2) or "") for r in cleaned
                          if r.get(_fk1) and r.get(_fk2)]
                _fuk = list(dict.fromkeys(_fkeys))
                _fexists = set()
                for _fi in range(0, len(_fuk), 200):
                    _fb = _fuk[_fi:_fi+200]
                    _fp = ",".join(["%s"] * len(_fb))
                    _fsoft = (" AND (deleted_at IS NULL OR deleted_at='')" if _ftbl == "orders" else "")
                    _frow = query("SELECT CONCAT(`%s`,'|',`%s`) k FROM `%s` WHERE channel=%%s%s "
                                  "AND CONCAT(`%s`,'|',`%s`) IN (%s)" % (_fk1, _fk2, _ftbl, _fsoft,
                                                                          _fk1, _fk2, _fp),
                                  [channel] + _fb)
                    for _r2 in _frow:
                        _fexists.add(_r2.get("k") or "")
                if _fexists:
                    cleaned = [r for r in cleaned if not (str(r.get(_fk1) or "") + "|" + str(r.get(_fk2) or "") in _fexists)]
            except Exception:
                pass
        w_success, w_failed, err_details = _write_rows(target, channel, conflict_mode, cleaned)
        success += w_success
        failed += w_failed
        # 2026-10-16 失败明细累积(跨窗口→done 回传前端 failed_details——曾窗口间丢弃无因可查)
        _fd = list(params.get("err_details") or [])
        for _e1 in (err_details or []):
            if len(_fd) < 50:
                _fd.append(_e1)
        offset += len(window)
        # 2026-10-16 窗口写后计数留痕(message 含 task_id——唯一不禁频)——窗口5 failed 1000 定位
        try:
            from db import execute as _ew
            _ew("INSERT INTO quality_logs(log_type, level, message, source) VALUES('perf_debug','info',%s,'cleansing')",
                ("[w%d] %s 窗写 success=%d failed=%d offset=%d/%d" % (offset // page_size, task_id, w_success, w_failed, offset, total)))
        except Exception:
            pass
        done = offset >= total
        if done:
            # 2026-10-16 done 步前置 finalizing 标记(防重跑: 锁3s会释放——重试重跑窗口5→唯一索引冲突→failed 覆盖——标记后重试不重跑)
            try:
                from db import execute as _ef
                _ef("UPDATE sync_tasks SET params=%s WHERE task_id=%s",
                    [json.dumps({**params, "finalizing": 1, "success": success, "failed": failed, "offset": offset, "total": total}, ensure_ascii=False), task_id])
            except Exception:
                pass
            # 最后步: 全量重清洗(联动/评估/汇总用)
            _st0 = _t.time()
            cleaned_all = _clean_rows(rows, mp)
            try_err('cleansing', 'done步 清洗耗时: %s' % round(_t.time() - _st0, 1))
            if target == "order":
                try:
                    from routes.suppliers import _norm_column_value as _ncol2
                    for x in cleaned_all:
                        if x.get("order_status"):
                            x["order_status"] = _ncol2(channel, "order_status", str(x["order_status"]))
                except Exception as _e:
                        try_err('cleansing', '静默降级', _e)
            elapsed = round(_t.time() - started, 1)
            # 2026-10-16 分离: 导入 done(写 done 计数——联动评估由独立 cleansing_eval 任务续跑分片——根治 done 步 80s 请求被杀)
            _steps_done = (params.get("steps") or []) + [{"name": "窗口 %d" % ((offset - len(window)) // page_size + 1), "elapsed": round(_t.time() - started, 1)}]
            execute("UPDATE sync_tasks SET status='done', params=%s, result=%s, updated_at=NOW() WHERE task_id=%s",
                    (json.dumps({**params, "offset": offset, "success": success, "failed": failed,
                                 "total": total, "steps": _steps_done, "cur_step": None}, ensure_ascii=False),
                     json.dumps({"result": {"target": target, "success": success, "failed": failed,
                                            "elapsed": elapsed,
                                            "message": "成功 %d 条, 跳过 %d 条 · 评估联动已创建独立任务" % (success, failed),
                                            "failed_details": _fd if _fd else [],
                                            "eval_task": task_id + "_eval"}}, ensure_ascii=False), task_id))
            # 2026-10-16 分离: 创建评估联动任务(独立续跑——分片评估不超时——导入 done 后自动触发)
            try:
                _ev_skus = []
                _ev_seen = set()
                # 2026-10-17 修复: 曾从 cleaned(最后窗)提取——分窗接力时只有末窗 SKU 漏前面窗(异步1100行 sku_list=100 实测);
                # 改从 cleaned_all(全量重清洗)提取——所有窗口 SKU 全覆盖
                for _ec in cleaned_all:
                    _esk = str(_ec.get("sku") or "").strip()
                    if _esk and _esk not in _ev_seen:
                        _ev_seen.add(_esk)
                        _ev_skus.append(_esk)
                from db import execute as _ee2
                _eval_task = task_id + "_eval"
                _ee2("INSERT INTO sync_tasks(task_id, task_type, status, params, result, channel) "
                     "VALUES(%s,'cleansing_eval','running',%s,'{}',%s) "
                     "ON DUPLICATE KEY UPDATE status='running', params=VALUES(params), updated_at=NOW()",
                     [_eval_task, json.dumps({"src_task": task_id, "target": target, "channel": channel,
                                              "mp": mp, "fname": fname, "offset": 0, "success": 0, "failed": 0,
                                              "sku_list": _ev_skus, "steps": [],
                                              "cur_step": {"name": "评估联动准备", "started": int(_t.time())}}, ensure_ascii=False), channel])
            except Exception as _e4:
                try_err('cleansing', '创建评估任务降级', _e4)
            from routes.analysis_cache import invalidate_all
            invalidate_all()
            return {"ok": True, "task_id": task_id, "status": "done",
                    "success": success, "failed": failed,
                    "message": "成功 %d 条, 跳过 %d 条 · 评估联动已创建独立任务" % (success, failed),
                    "failed_details": _fd if _fd else []}
        # 未完 → 更新 offset/累积, running(2026-10-15 每窗 steps 耗时 + cur_step 进度中——前端 TaskCard 可视化)
        _sq = round(_t.time() - started, 1)
        _steps = (params.get("steps") or []) + [{"name": "窗口 %d" % ((offset - len(window)) // page_size + 1), "elapsed": _sq}]
        _cur = {"name": "窗口 %d" % (offset // page_size + 1), "started": int(_t.time())}
        execute("UPDATE sync_tasks SET params=%s, updated_at=NOW() WHERE task_id=%s",
                (json.dumps({**params, "offset": offset, "success": success, "failed": failed,
                             "total": total, "steps": _steps, "cur_step": _cur}, ensure_ascii=False), task_id))
        return {"ok": True, "task_id": task_id, "status": "running",
                "page": offset // page_size, "total": total}
    except Exception as e:
        import traceback as _tb
        execute("UPDATE sync_tasks SET status='error', result=%s, updated_at=NOW() WHERE task_id=%s",
                (json.dumps({"error": str(e)[:400], "tb": _tb.format_exc()[-1200:]},
                            ensure_ascii=False), task_id))
        try_err("cleansing", "导入续跑失败", e)
        return {"ok": True, "task_id": task_id, "status": "error", "error": str(e)[:200]}


def _b64(code):
    import base64 as _b
    return _b.b64encode(code).decode("ascii")


@router.get("/cleansing/task/{task_id}")
@traced
def cleansing_task(task_id: str):
    row = one("SELECT status, result FROM sync_tasks WHERE task_id=%s", [task_id])
    if not row:
        return {"status": "not_found"}
    status = row.get("status") or "pending"
    result = {}
    try:
        result = json.loads(row.get("result") or "{}")
    except Exception as _e:
            try_err('cleansing', '静默降级', _e)
    out = {"status": status}
    if status == "done":
        out["result"] = result.get("result") or result
    elif status == "error":
        out["error"] = result.get("error") or "清洗失败"
    return out


# ── 模板 CRUD ───────────────────────────────────────────────────────────
@router.delete("/cleansing/templates/{tid}")
@traced
def cleansing_templates_delete(tid: int):
    """删除清洗模板"""
    execute("DELETE FROM cleansing_templates WHERE id=%s", [tid])
    return ok({"message": "已删除"})


@router.get("/cleansing/templates")
@traced
def cleansing_templates():
    rows = query("SELECT id, name, doc_type, mapping, cf_meta FROM cleansing_templates ORDER BY id DESC LIMIT 100")
    return ok(rows)


@router.get("/cleansing/eval-status")
@traced
def cleansing_eval_status(task_id: str = "", page_size: int = 1000):
    """2026-10-16 评估联动独立任务——分窗接力(每请求 ≤90s 安全——SKU 清单驱动——不重读源)
    读 sku_list → 处理 offset 起 page_size 个 SKU(评估+联动该片——幂等 existing 去重+事务联动) → offset 推进 → done"""
    import time as _t
    row = one("SELECT status, params, result FROM sync_tasks WHERE task_id=%s", [task_id])
    if not row:
        return {"ok": False, "error": "任务不存在"}
    status = row.get("status") or ""
    if status in ("done", "error"):
        _ri = (json.loads(row.get("result") or "{}") or {}).get("result") or {}
        return {"ok": True, "task_id": task_id, "status": status, "success": _ri.get("success"),
                "failed": _ri.get("failed"), "error": _ri.get("error"),
                "rules_evaluated": _ri.get("rules_evaluated"), "inventory_adjusted": _ri.get("inventory_adjusted")}
    params = json.loads(row.get("params") or "{}")
    channel = params.get("channel") or "jd"
    target = params.get("target") or "order"
    sku_list = params.get("sku_list") or []
    offset = int(params.get("offset") or 0)
    started = _t.time()
    try:
        # 抢占锁(防并发续跑——同导入)
        _locked = execute("UPDATE sync_tasks SET updated_at=NOW() WHERE task_id=%s "
                          "AND (updated_at IS NULL OR updated_at < DATE_SUB(NOW(), INTERVAL 3 SECOND))",
                          [task_id])
        if not _locked:
            return {"ok": True, "task_id": task_id, "status": "running"}
        _batch = sku_list[offset:offset + page_size]
        if not _batch:
            # 全部分片完成 → done(累计 result)
            _rj = json.loads(row.get("result") or "{}") or {}
            _rri = _rj.get("result") or {}
            execute("UPDATE sync_tasks SET status='done', params=%s, updated_at=NOW() WHERE task_id=%s",
                    (json.dumps({**params, "cur_step": None}, ensure_ascii=False), task_id))
            from routes.analysis_cache import invalidate_all
            invalidate_all()
            return {"ok": True, "task_id": task_id, "status": "done",
                    "rules_evaluated": _rri.get("rules_evaluated"), "inventory_adjusted": _rri.get("inventory_adjusted")}
        _total = len(sku_list)
        _seg_no = offset // page_size + 1
        _seg_cnt = (_total + page_size - 1) // page_size
        # 该片 SKU: 评估(order.created——库存上下文)+ 联动(销售扣减——deltas 只算该片)
        from db import query as _q5, execute as _e5, transaction as _txn5
        from routes.cleanse_write import _adjust_inventory
        from core.rules import evaluate_many, load_rules_for
        from collections import defaultdict
        inv_map = {}
        for _i in range(0, len(_batch), 500):
            _ph = ",".join(["%s"] * len(_batch[_i:_i + 500]))
            for r in _q5("SELECT sku, MAX(available_qty) AS avail FROM inventory "
                         "WHERE channel=%s AND sku IN (%s) GROUP BY sku" % ("%s", _ph),
                         [channel] + _batch[_i:_i + 500]):
                inv_map[r.get("sku")] = int(r.get("avail") or 0)
        # 从源文件取该片 SKU 的订单行(quantity 聚合——评估 ctx 行级?——SKU 级聚合: order_qty=该 SKU 该片总量)
        src_task = params.get("src_task") or ""
        import base64 as _b64e
        fr = one("SELECT content FROM export_files WHERE filename=%s", ["clean_" + src_task])
        c = fr.get("content") if fr else None
        _data = _b64e.b64decode(c[7:]) if isinstance(c, str) and c.startswith("base64:") else (c.encode("utf-8") if isinstance(c, str) else (bytes(c) if c is not None else b""))
        _mpd = params.get("mp") or {}
        _mpd = _mpd if isinstance(_mpd, dict) else json.loads(_mpd or "{}")
        headers, rows = _parse_table(_data, params.get("fname") or "import.csv")
        _cleaned = _clean_rows(rows, _mpd)
        # 该片 SKU 行级评估 ctx(与导入 done 步同口径——每行 ctx——order.created 规则行级判定)
        # 2026-10-29 超卖口径修复: 原 ctx 只给 order.quantity=本次导入量(单笔) →
        # 累计已超卖(如 1095 件 vs 库存 9)但单笔都小于库存时**永不报警**(系统性漏报)
        # → 补 order.total_qty = 该 SKU 累计未删订单量(与 close_oversell 判定口径一致)
        _cum = {}
        if _batch:
            _ph2 = ",".join(["%s"] * len(_batch[:500]))
            for _r in _q5("SELECT sku, SUM(quantity) AS q FROM orders WHERE channel=%s "
                          "AND (deleted_at IS NULL OR deleted_at='') AND sku IN (%s) "
                          "GROUP BY sku" % ("%s", _ph2), [channel] + _batch[:500]):
                _cum[_r.get("sku")] = int(_r.get("q") or 0)
        # 2026-10-29 SKU 级去重评估(数学等价): 原每行一 ctx, 同 SKU 多行重复评估
        #   等价性: "存在某行 quantity > X" ⟺ max_quantity > X; "< X" ⟺ min_quantity < X
        #   → 提供 max/min/total_qty/line_count 四个字段, 任何单调条件都可精确表达(语义不漂移)
        _grp = {}
        for c in _cleaned:
            sku = str(c.get("sku") or "").strip()
            if sku not in set(_batch):
                continue
            _q = int(c.get("quantity") or 0)
            _g = _grp.setdefault(sku, {"max": _q, "min": _q, "n": 0})
            _g["max"] = max(_g["max"], _q)
            _g["min"] = min(_g["min"], _q)
            _g["n"] += 1
        # 2026-10-29 入库/出库触发新事件(此前完全不评估 —— 到货/出库后库存骤降无法即时预警)
        # 约束(四维): 仅本次导入 SKU(不全量) + try/except 降级(不阻塞) + 异步任务内执行(不拖慢导入)
        _flow_evt = {"inbound": "inbound.created", "outbound": "outbound.created"}.get(
            str(params.get("target") or ""))
        flow_ctxs = []
        if _flow_evt:
            _fgrp = {}
            for _fc in _cleaned:
                _fsku = str(_fc.get("sku") or "").strip()
                if not _fsku or _fsku not in set(_batch):
                    continue
                _fq = int(_fc.get("quantity") or 0)
                _fg = _fgrp.setdefault(_fsku, {"qty": 0, "wh": "", "date": "", "exp": ""})
                _fg["qty"] += _fq
                _fg["wh"] = _fg["wh"] or str(_fc.get("warehouse") or "")
                _fg["date"] = _fg["date"] or str(
                    _fc.get("inbound_date") or _fc.get("outbound_date") or "")
                _fg["exp"] = _fg["exp"] or str(_fc.get("exp_date") or "")
            for _fs, _fg in _fgrp.items():
                flow_ctxs.append({"sku": _fs, "channel": channel,
                                  "flow": {"quantity": _fg["qty"], "warehouse": _fg["wh"],
                                           "date": _fg["date"]},
                                  "inbound": {"quantity": _fg["qty"], "warehouse": _fg["wh"],
                                              "date": _fg["date"]},
                                  "outbound": {"quantity": _fg["qty"], "warehouse": _fg["wh"],
                                               "date": _fg["date"]},
                                  "available_stock": inv_map.get(_fs, 0),
                                  "inv": {"available_qty": inv_map.get(_fs, 0),
                                          "safety_qty": 0, "in_transit_qty": 0}})
        o_ctxs = []
        for sku, _g in _grp.items():
            o_ctxs.append({"sku": sku, "channel": channel,
                           "order": {"quantity": _g["max"],
                                     "max_quantity": _g["max"],
                                     "min_quantity": _g["min"],
                                     "line_count": _g["n"],
                                     "total_qty": _cum.get(sku, _g["max"])},
                           "order_qty": _cum.get(sku, _g["max"]),
                           "available_stock": inv_map.get(sku, 0),
                           "inv": {"available_qty": inv_map.get(sku, 0), "safety_qty": 0, "in_transit_qty": 0}})
        evaluated = 0
        if flow_ctxs and _flow_evt:
            try:
                _frules = load_rules_for(_flow_evt, channel)
                if _frules:
                    evaluate_many(_flow_evt, flow_ctxs, channel, _frules)
                    evaluated += len(flow_ctxs)
            except Exception as _fe:
                try_err('cleansing', '出入库评估降级', _fe)
        if o_ctxs:
            evaluate_many("order.created", o_ctxs, channel, load_rules_for("order.created", channel))
            evaluated = len(o_ctxs)
        # 联动(该片 SKU 销售扣减——短事务原子)
        deltas = defaultdict(int)
        for c in _cleaned:
            sku = str(c.get("sku") or "").strip()
            if sku not in set(_batch):
                continue
            _qty = int(c.get("quantity") or 0)
            _wh = str(c.get("warehouse") or "")
            _ds = ((_mpd or {}).get("_meta") or {}).get("data_source", "")
            _d = _qty if _ds == "jd_po" else -_qty
            if _d:
                deltas[(sku, _wh, "platform")] += _d
        adjusted = 0
        try:
            with _txn5():
                if deltas:
                    adjusted, _ = _adjust_inventory(channel, {k: v for k, v in deltas.items()})
        except Exception as _e5b:
            try_err('cleansing', '联动事务异常(回滚可重试)', _e5b)
            return {"ok": True, "task_id": task_id, "status": "running"}
        # 推进 offset + 累积 result
        _new_off = offset + len(_batch)
        _rj = json.loads(row.get("result") or "{}") or {}
        _rri = _rj.get("result") or {}
        _cum_rules = int(_rri.get("rules_evaluated") or 0) + evaluated
        _cum_inv = int(_rri.get("inventory_adjusted") or 0) + adjusted
        _steps = (params.get("steps") or []) + [{"name": "评估分片 %d/%d" % (_seg_no, _seg_cnt), "elapsed": round(_t.time() - started, 1)}]
        _done = _new_off >= _total
        _status = "done" if _done else "running"
        execute("UPDATE sync_tasks SET status=%s, params=%s, result=%s, updated_at=NOW() WHERE task_id=%s",
                (_status,
                 json.dumps({**params, "offset": _new_off, "steps": _steps,
                             "cur_step": None if _done else {"name": "评估分片 %d/%d" % (_seg_no + 1, _seg_cnt), "started": int(_t.time())}}, ensure_ascii=False),
                 json.dumps({"result": {"target": target, "success": _total, "failed": 0,
                                        "rules_evaluated": _cum_rules, "inventory_adjusted": _cum_inv,
                                        "message": "评估 %d SKU, 联动 %d" % (_cum_rules, _cum_inv)}}, ensure_ascii=False), task_id))
        if _done:
            from routes.analysis_cache import invalidate_all
            invalidate_all()
        return {"ok": True, "task_id": task_id, "status": _status, "page": _seg_no,
                "rules_evaluated": _cum_rules, "inventory_adjusted": _cum_inv}
    except Exception as _e6:
        try_err('cleansing', '评估任务异常(可重试)', _e6)
        return {"ok": True, "task_id": task_id, "status": "running"}


@router.post("/cleansing/templates")
@traced
async def cleansing_templates_save(request: Request):
    d = {}
    try:
        d = await request.json()
    except Exception as _e:
            try_err('cleansing', '静默降级', _e)
    name = d.get("name") or ""
    doc_type = d.get("doc_type") or "order"
    if not name:
        return fail("缺少模板名称")
    mapping = d.get("mapping")
    if isinstance(mapping, dict):
        mapping = json.dumps(mapping, ensure_ascii=False)
    # 2026-10-27 根治: 模板内嵌自定义字段定义(跨设备/清缓存可还原字段名)
    cf_meta = d.get("cf_meta") or []
    if isinstance(cf_meta, (list, dict)):
        cf_meta = json.dumps(cf_meta, ensure_ascii=False)
    # 同名模板更新覆盖(name + doc_type 同时存在 → UPDATE; 否则 INSERT——防同名重复)
    # 2026-10-27 隔离修复: 曾只按 name 查——不同导入类型同名模板互相覆盖(订单/库存同名=串类型)
    existing = one("SELECT id FROM cleansing_templates WHERE name=%s AND doc_type=%s",
                   [name, doc_type])
    if existing:
        execute("UPDATE cleansing_templates SET mapping=%s, doc_type=%s, cf_meta=%s, updated_at=NOW() "
                "WHERE id=%s",
                (mapping or "{}", doc_type, cf_meta or "[]", existing.get("id")))
    else:
        execute("INSERT INTO cleansing_templates(name, doc_type, mapping, cf_meta) "
                "VALUES(%s,%s,%s,%s)",
                (name, doc_type, mapping or "{}", cf_meta or "[]"))
    return ok({"message": "模板已保存"})


# ── 写入逻辑(按目标类型) ────────────────────────────────────────────────





